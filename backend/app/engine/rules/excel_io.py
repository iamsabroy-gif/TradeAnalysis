"""
Excel Import / Export service for the Dynamic Configuration Rule Engine.
Reads and writes `Rules_Config.xlsx` using openpyxl.
Strictly maps to Docs/rule-engine-implementation.md §1.
"""

import io
import re
from typing import Any, Dict, List, Optional, Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.app.models.enums import Phase2Sector
from backend.app.engine.rules.config import (
    Phase1RuleConfig,
    Phase2MatrixConfig,
    RulesConfiguration,
    SectorKeywordItem,
    SectorThresholdConfig,
    create_default_rules_configuration,
)


# Styles for Excel Generation
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="64748B")
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")


def _clean_num(val: Any) -> Optional[float]:
    """Parses float from number, string, or percentage representation."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "")
    if not s:
        return None
    # Strip trailing multiplier 'x' (e.g. '3.0x' -> 3.0)
    if s.lower().endswith("x"):
        s = s[:-1].strip()
    # Strip percentage '%' (e.g. '10%' -> 10.0)
    if s.endswith("%"):
        try:
            return float(s[:-1].strip())
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None


def generate_default_rules_workbook() -> bytes:
    """
    Generates the canonical 3-sheet `Rules_Config.xlsx` workbook template.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    config = create_default_rules_configuration()

    # ----------------------------------------------------
    # Sheet 1: Phase1_Thresholds
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="Phase1_Thresholds")
    ws1.views.sheetView[0].showGridLines = True

    headers1 = ["Check #", "Metric", "Fail_Threshold", "Operator", "Sector_Aware?", "Note"]
    ws1.append(headers1)

    rows1 = [
        [2, "Promoter_Pledge", f"{config.phase1.promoter_pledge_fail_pct}%", ">", "No", "Pledge % of promoter holding triggering FAIL"],
        [2, "Promoter_Pledge_Low_Holding_Floor", f"{config.phase1.promoter_pledge_low_holding_floor_pct}%", "<", "No", "Promoter holding floor triggering tighter pledge rule"],
        [2, "Promoter_Pledge_Low_Holding_Fail", f"{config.phase1.promoter_pledge_low_holding_fail_pct}%", ">", "No", "Tighter pledge ceiling when promoter holding < 35%"],
        [2, "Promoter_Pledge_Total_Shares", f"{config.phase1.promoter_pledge_total_shares_fail_pct}%", ">", "No", "Pledged shares as % of total shares triggering FAIL"],
        [3, "RPT_Sales_Purchases", f"{config.phase1.rpt_sales_purchases_fail_pct}%", ">", "No", "RPT sales + purchases as % of revenue"],
        [4, "Contingent_Liabilities", f"{config.phase1.contingent_liabilities_net_worth_fail_pct}%", ">", "Yes", "Total contingent liabilities as % of net worth"],
        [4, "Litigation_Claims", f"{config.phase1.litigation_claims_net_worth_fail_pct}%", ">", "No", "Tax & litigation claims as % of net worth"],
        [5, "CFO_PAT_Ratio", f"{config.phase1.cfo_pat_ratio_fail_floor:.2f}", "<", "No", "5-year cumulative CFO / PAT conversion floor"],
        [6, "Legal_Fee_Surge", f"{config.phase1.legal_fee_surge_fail_pct}%", ">", "No", "YoY legal fee surge triggering FAIL if unexplained"],
        [6, "Legal_Audit_Fee_Multiplier", f"{config.phase1.legal_to_audit_fee_multiplier:.1f}x", ">", "No", "Ratio of legal fees to audit fees triggering FAIL"],
    ]

    for r_idx, row in enumerate(rows1, start=2):
        ws1.append(row)

    # Style Sheet 1
    for col in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx in range(2, len(rows1) + 2):
        is_alt = (r_idx % 2 == 0)
        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=r_idx, column=col_idx)
            cell.font = ROW_FONT
            cell.border = THIN_BORDER
            if is_alt:
                cell.fill = ALT_FILL
            if col_idx in [1, 3, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                cell.font = NOTE_FONT

    # ----------------------------------------------------
    # Sheet 2: Phase2_Matrix
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Phase2_Matrix")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = ["Sector Profile", "Metric", "Pass_Threshold", "Fail_Threshold", "Operator", "Unit", "Note"]
    ws2.append(headers2)

    rows2 = []
    for sector_key, thresh in config.phase2_matrix.items():
        rows2.extend([
            [thresh.sector.value, "RoCE", f"{thresh.roce_pass_floor}%", f"{thresh.roce_fail_ceiling}%", "Range", "%", "Return on Capital Employed (5-yr median)"],
            [thresh.sector.value, "NetDebt_EBITDA", f"{thresh.net_debt_ebitda_pass_ceiling:.1f}x", f"{thresh.net_debt_ebitda_fail_floor:.1f}x", "Range", "x", "Net Debt to EBITDA (pass <= ceiling, fail > floor)"],
            [thresh.sector.value, "Interest_Coverage", f"{thresh.interest_coverage_pass_floor:.1f}x", f"{thresh.interest_coverage_fail_ceiling:.1f}x", "Range", "x", "EBIT to Finance Cost coverage floor"],
            [thresh.sector.value, "Max_ST_Debt_Pct", f"{thresh.max_st_debt_concern_pct}%", f"{thresh.max_st_debt_concern_pct}%", ">", "%", "Short-term debt % of total debt concern threshold"],
            [thresh.sector.value, "CCC_Deterioration_Days", f"{thresh.ccc_deterioration_fail_days:.0f}", f"{thresh.ccc_deterioration_fail_days:.0f}", ">", "days", "3-year rise in Cash Conversion Cycle"],
            [thresh.sector.value, "Max_Receivable_Days", f"{thresh.max_receivable_days_critical_cap:.0f}", f"{thresh.max_receivable_days_critical_cap:.0f}", ">", "days", "Critical receivable days ceiling"],
        ])

    for row in rows2:
        ws2.append(row)

    # Style Sheet 2
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx in range(2, len(rows2) + 2):
        is_alt = (r_idx % 2 == 0)
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.font = ROW_FONT
            cell.border = THIN_BORDER
            if is_alt:
                cell.fill = ALT_FILL
            if col_idx in [1, 3, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 7:
                cell.font = NOTE_FONT

    # ----------------------------------------------------
    # Sheet 3: Sector_Mapping
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="Sector_Mapping")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = ["Industry Keyword", "Sector Profile", "Notes"]
    ws3.append(headers3)

    rows3 = []
    for item in config.sector_mappings:
        rows3.append([item.keyword, item.sector.value, item.notes or ""])

    for row in rows3:
        ws3.append(row)

    # Style Sheet 3
    for col in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx in range(2, len(rows3) + 2):
        is_alt = (r_idx % 2 == 0)
        for col_idx in range(1, len(headers3) + 1):
            cell = ws3.cell(row=r_idx, column=col_idx)
            cell.font = ROW_FONT
            cell.border = THIN_BORDER
            if is_alt:
                cell.fill = ALT_FILL
            if col_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 3:
                cell.font = NOTE_FONT

    # Auto-fit column widths for all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_rules_config_workbook(
    file_bytes_or_stream: Union[bytes, io.BytesIO, str]
) -> RulesConfiguration:
    """
    Parses an uploaded `Rules_Config.xlsx` workbook into a strongly typed RulesConfiguration.
    Strictly handles all 3 sheets per Docs/rule-engine-implementation.md §1.1.
    """
    if isinstance(file_bytes_or_stream, bytes):
        stream = io.BytesIO(file_bytes_or_stream)
    elif isinstance(file_bytes_or_stream, str):
        stream = open(file_bytes_or_stream, "rb")
    else:
        stream = file_bytes_or_stream

    wb = openpyxl.load_workbook(stream, data_only=True)

    # Verify sheet presence
    required_sheets = ["Phase1_Thresholds", "Phase2_Matrix", "Sector_Mapping"]
    for s in required_sheets:
        if s not in wb.sheetnames:
            raise ValueError(f"Invalid Rules_Config.xlsx: Missing required sheet '{s}'")

    # 1. Parse Sheet 1: Phase1_Thresholds
    ws1 = wb["Phase1_Thresholds"]
    p1_dict: Dict[str, float] = {}

    for row in ws1.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        metric_name = str(row[1]).strip()
        val = _clean_num(row[2])
        if val is not None:
            p1_dict[metric_name.lower()] = val

    # Build Phase1RuleConfig with fallback defaults if metric omitted
    p1_cfg = Phase1RuleConfig(
        promoter_pledge_fail_pct=p1_dict.get("promoter_pledge", 10.0),
        promoter_pledge_low_holding_floor_pct=p1_dict.get("promoter_pledge_low_holding_floor", 35.0),
        promoter_pledge_low_holding_fail_pct=p1_dict.get("promoter_pledge_low_holding_fail", 5.0),
        promoter_pledge_total_shares_fail_pct=p1_dict.get("promoter_pledge_total_shares", 5.0),
        rpt_sales_purchases_fail_pct=p1_dict.get("rpt_sales_purchases", 10.0),
        contingent_liabilities_net_worth_fail_pct=p1_dict.get("contingent_liabilities", 15.0),
        litigation_claims_net_worth_fail_pct=p1_dict.get("litigation_claims", 10.0),
        cfo_pat_ratio_fail_floor=p1_dict.get("cfo_pat_ratio", 0.80),
        legal_fee_surge_fail_pct=p1_dict.get("legal_fee_surge", 100.0),
        legal_to_audit_fee_multiplier=p1_dict.get("legal_audit_fee_multiplier", 3.0),
    )

    # 2. Parse Sheet 2: Phase2_Matrix
    ws2 = wb["Phase2_Matrix"]
    matrix_raw: Dict[str, Dict[str, Dict[str, float]]] = {}

    def _normalize_sector(sec_str: str) -> Optional[Phase2Sector]:
        norm = sec_str.strip().upper().replace(" ", "_").replace("-", "_")
        for s in Phase2Sector:
            if s.value == norm or s.name == norm:
                return s
        if "ASSET" in norm:
            return Phase2Sector.ASSET_LIGHT
        if "CAP" in norm:
            return Phase2Sector.CAP_INTENSIVE
        if "REGULAT" in norm or "INFRA" in norm or "UTILITY" in norm:
            return Phase2Sector.REGULATED_INFRA
        if "STAND" in norm:
            return Phase2Sector.STANDARD
        return None

    for row in ws2.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        sec_enum = _normalize_sector(str(row[0]))
        if not sec_enum:
            continue
        metric = str(row[1]).strip().lower()
        pass_val = _clean_num(row[2])
        fail_val = _clean_num(row[3])

        sec_key = sec_enum.value
        if sec_key not in matrix_raw:
            matrix_raw[sec_key] = {}
        matrix_raw[sec_key][metric] = {
            "pass": pass_val if pass_val is not None else 0.0,
            "fail": fail_val if fail_val is not None else 0.0,
        }

    # Populate Phase2 Matrix with defaults as fallback for any missing metric
    default_cfg = create_default_rules_configuration()
    matrix_final: Dict[str, SectorThresholdConfig] = {}

    for sec in Phase2Sector:
        sec_key = sec.value
        def_sec = default_cfg.phase2_matrix.get(sec_key)
        sec_data = matrix_raw.get(sec_key, {})

        roce_d = sec_data.get("roce", {})
        debt_d = sec_data.get("netdebt_ebitda", {})
        int_d = sec_data.get("interest_coverage", {})
        st_d = sec_data.get("max_st_debt_pct", {})
        ccc_d = sec_data.get("ccc_deterioration_days", {})
        rec_d = sec_data.get("max_receivable_days", {})

        matrix_final[sec_key] = SectorThresholdConfig(
            sector=sec,
            roce_pass_floor=roce_d.get("pass", def_sec.roce_pass_floor if def_sec else 15.0),
            roce_fail_ceiling=roce_d.get("fail", def_sec.roce_fail_ceiling if def_sec else 10.0),
            net_debt_ebitda_pass_ceiling=debt_d.get("pass", def_sec.net_debt_ebitda_pass_ceiling if def_sec else 2.0),
            net_debt_ebitda_fail_floor=debt_d.get("fail", def_sec.net_debt_ebitda_fail_floor if def_sec else 3.0),
            interest_coverage_pass_floor=int_d.get("pass", def_sec.interest_coverage_pass_floor if def_sec else 4.0),
            interest_coverage_fail_ceiling=int_d.get("fail", def_sec.interest_coverage_fail_ceiling if def_sec else 2.0),
            max_st_debt_concern_pct=st_d.get("fail", def_sec.max_st_debt_concern_pct if def_sec else 20.0),
            ccc_deterioration_fail_days=ccc_d.get("fail", def_sec.ccc_deterioration_fail_days if def_sec else 60.0),
            max_receivable_days_critical_cap=rec_d.get("fail", def_sec.max_receivable_days_critical_cap if def_sec else 90.0),
        )

    # 3. Parse Sheet 3: Sector_Mapping
    ws3 = wb["Sector_Mapping"]
    mappings: List[SectorKeywordItem] = []

    for row in ws3.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        kw = str(row[0]).strip()
        sec_enum = _normalize_sector(str(row[1]))
        notes = str(row[2]).strip() if len(row) > 2 and row[2] else None
        if kw and sec_enum:
            mappings.append(SectorKeywordItem(keyword=kw, sector=sec_enum, notes=notes))

    if not mappings:
        mappings = default_cfg.sector_mappings

    return RulesConfiguration(
        version="1.0.0",
        source="EXCEL_UPLOAD",
        description="User Custom Rules Configuration loaded via Excel upload",
        phase1=p1_cfg,
        phase2_matrix=matrix_final,
        sector_mappings=mappings,
    )
