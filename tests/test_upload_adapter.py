"""
Unit tests for Workbook Upload Adapter and intake security.
Strictly maps to Phase1-PhaseC-Scraper-Implementation-Plan.md §5A & §7.
"""

import io
import openpyxl
import pytest

from backend.app.acquisition.uploads.intake import (
    FileValidationError,
    validate_and_hash_upload,
)
from backend.app.acquisition.uploads.template import generate_workbook_template_bytes
from backend.app.acquisition.uploads.workbook import WorkbookUploadAdapter
from backend.app.models.enums import AuditOpinion, Confidence, ExtractionMethod


def test_template_generator():
    template_bytes = generate_workbook_template_bytes()
    assert len(template_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == ["field_name", "value", "source", "period", "basis", "notes"]

    # Check some key fields exist in column A
    col_a = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "promoter_holding_pct_of_company" in col_a
    assert "contingent_liabilities" in col_a
    assert "audit_opinion" in col_a


def test_workbook_adapter_csv_valid():
    adapter = WorkbookUploadAdapter(uploader="test_analyst")
    csv_content = """field_name,value,source,period,basis,notes
promoter_holding_pct_of_company,60.0,FY24 AR p. 45,FY24,NOT_APPLICABLE,
audit_opinion,CLEAN,FY24 Independent Auditor Report,FY24,NOT_APPLICABLE,
cfo_last_5y,"100, 110, 120, 130, 140",AR Cash Flow,FY20-FY24,CONSOLIDATED,
contingent_liabilities,120.0,AR Note 31,FY24,CONSOLIDATED,
"""
    result = adapter.parse_bytes(csv_content.encode("utf-8"), filename="test_workbook.csv")
    assert len(result.errors) == 0
    assert len(result.fields) == 4

    field_map = {f.field_name: f for f in result.fields}
    assert field_map["promoter_holding_pct_of_company"].value == 60.0
    assert field_map["promoter_holding_pct_of_company"].confidence == Confidence.MANUAL
    assert field_map["audit_opinion"].value == AuditOpinion.CLEAN
    assert field_map["cfo_last_5y"].value == [100.0, 110.0, 120.0, 130.0, 140.0]


def test_workbook_adapter_validation_rejections():
    adapter = WorkbookUploadAdapter(uploader="test_analyst")
    bad_csv = """field_name,value,source,period,basis,notes
unknown_field,100,Source A,FY24,NOT_APPLICABLE,
promoter_holding_pct_of_company,50.0,,FY24,NOT_APPLICABLE,
cfo_last_5y,"100, 200, 300",AR Cash Flow,FY24,CONSOLIDATED,
"""
    result = adapter.parse_bytes(bad_csv.encode("utf-8"), filename="bad.csv")
    assert len(result.errors) == 3

    err_fields = [e.field_name for e in result.errors]
    assert "unknown_field" in err_fields
    assert "promoter_holding_pct_of_company" in err_fields
    assert "cfo_last_5y" in err_fields


def test_intake_validation_security():
    # Unsupported extension
    with pytest.raises(FileValidationError, match="Unsupported file extension"):
        validate_and_hash_upload("malicious.exe", b"bad content")

    # Size cap
    oversized = b"x" * (3 * 1024 * 1024)  # 3 MB
    with pytest.raises(FileValidationError, match="exceeds maximum allowed size"):
        validate_and_hash_upload("sheet.xlsx", oversized)
